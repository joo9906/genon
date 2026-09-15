import os
import sys
import uuid
import json
import mimetypes
import tempfile
import re
import subprocess
import importlib 
from urllib import request
from typing import Dict, Any, Optional

def ensure_packages():
    packages = {"docx": "python-docx"}
    for import_name, install_name in packages.items():
        try:
            importlib.import_module(import_name)
        except ImportError:
            print(f"패키지 설치 중: {install_name}...")
            subprocess.check_call([
                sys.executable, "-m", "pip", "install", install_name, "--upgrade"
            ])

ensure_packages()

try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    Document = None


class GenOSMinIOService:
    def __init__(self, url: str = None):
        self.url = url or "http://llmops-cdn-api-service:8080/minio/upload/temp"
        self.hostname = "https://genos.genon.ai"

    def _build_multipart_body(self, file_path: str):
        boundary = "----WebKitFormBoundary" + uuid.uuid4().hex
        crlf = "\r\n"
        filename = os.path.basename(file_path)
        with open(file_path, "rb") as f:
            file_data = f.read()
        
        hostname_field = f"--{boundary}{crlf}Content-Disposition: form-data; name=\"hostname\"{crlf}{crlf}{self.hostname}{crlf}".encode("utf-8")
        content_type = mimetypes.guess_type(filename)[0] or "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        file_field = (f"--{boundary}{crlf}Content-Disposition: form-data; name=\"file\"; filename=\"{filename}\"{crlf}Content-Type: {content_type}{crlf}{crlf}").encode("utf-8") + file_data + crlf.encode("utf-8")
        body = hostname_field + file_field + f"--{boundary}--{crlf}".encode("utf-8")
        headers = {"Content-Type": f"multipart/form-data; boundary={boundary}", "Content-Length": str(len(body))}
        return body, headers

    def upload_file_and_get_url(self, file_path: str) -> str:
        body, headers = self._build_multipart_body(file_path)
        req = request.Request(self.url, data=body, headers=headers, method="POST")
        with request.urlopen(req) as resp:
            return json.loads(resp.read())["data"]["presigned_url"]


class DirectWordDocumentGenerator:
    """표지 없이 첫 장부터 제목과 본문이 바로 나오는 Word 생성 클래스"""

    @staticmethod
    def add_styled_text(paragraph, text: str):
        """마크다운 **굵게** 처리 엔진"""
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part.replace('**', ''))
                run.bold = True
            else:
                paragraph.add_run(part)

    def generate(self, config: Dict[str, Any]) -> Document:
        doc = Document()
        
        # --- [변경 지점 1] 최상단 대제목 배치 (표지 생략) ---
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.space_before = Pt(10)
        title_p.paragraph_format.space_after = Pt(24) # 제목 아래 여백
        
        title_run = title_p.add_run(config.get("report_title", "보고서"))
        title_run.font.size = Pt(17) # 크게 강조
        title_run.bold = True
        
        # --- 본문 렌더링 시작 ---
        content = config.get("content", "")
        processed_content = re.sub(r'</?br\s*/?>', ' ', content, flags=re.IGNORECASE)
        lines = processed_content.split('\n')
        
        table_data = []
        
        def flush_table():
            nonlocal table_data
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.style = config.get("table_style", "Table Grid")
                for i, row_cells in enumerate(table_data):
                    for j, cell_text in enumerate(row_cells):
                        self.add_styled_text(table.cell(i, j).paragraphs[0], cell_text)
                table_data = []
                doc.add_paragraph()

        for line in lines:
            clean_line = line.strip()
            if not clean_line:
                flush_table()
                doc.add_paragraph()
                continue

            # 마크다운 표(|) 판별
            if '|' in clean_line:
                if re.match(r'^[\s|:-]+$', clean_line): 
                    continue
                cells = [c.strip() for c in clean_line.split('|') if c.strip()]
                if cells: 
                    table_data.append(cells)
                continue
            
            flush_table()

            # 마크다운 헤더(#) 판별
            if clean_line.startswith('#'):
                level = 1
                if clean_line.startswith('###'): level = 3
                elif clean_line.startswith('##'): level = 2
                
                # 만약 content 내부 첫 라인에 제목이 중복문자열로 들어와도 처리 가능하도록 샵(#) 제거
                doc.add_heading(clean_line.replace('#', '').strip(), level=level)
                continue

            # 일반 본문
            p = doc.add_paragraph()
            self.add_styled_text(p, clean_line)

        flush_table()
        return doc


# MCP 툴 인터페이스
@mcp.tool()
async def generate_word_report_general_version(report_title: str, text: str) -> str:
    """
    표지 없이 첫 페이지 상단에 대제목을 배치하고, 바로 본문(마크다운) 내용이 이어지는 
    깔끔한 단일 형태의 Word 보고서(.docx)를 생성하여 업로드합니다.
    
    Args:
        report_title (str): 문서 최상단에 위치할 보고서 대제목
        text (str) : 작성된 마크다운 보고서 본문 내용 (content)
    """
    if Document is None:
        return "에러: 서버에 'python-docx' 패키지가 설치되어 있지 않습니다."

    try:
        # 데이터 매핑
        report_config = {
            "report_title": report_title,
            "content": text,
            "table_style": "Table Grid"
        }
        
        # 문서 빌드 (표지 없이 바로 생성)
        builder = DirectWordDocumentGenerator()
        doc = builder.generate(report_config)

        # 임시 파일 저장 및 업로드
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            doc.save(tmp.name)
            tmp_path = tmp.name

        service = GenOSMinIOService()
        presigned_url = service.upload_file_and_get_url(tmp_path)
        
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)
        
        return presigned_url

    except Exception as e:
        return f"보고서 생성 중 오류 발생: {str(e)}"